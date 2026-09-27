"""Kalibrasyon karşılaştırması için 2. tur baskı testi etiketleri.

A/B/C hepsi AYNI QR modül fiziksel boyutunda (5x5cm) basılsın diye — B/C'nin
kenar yaması şeridi QR'ın kendi boyutunu KüçÜLTMESİN, bunun yerine QR'ın
DIŞINA taşsın. Bu yüzden B/C, docx'e A'dan büyük bir fiziksel boyutta
gömülüyor (canvas piksel oranıyla orantılı) — QR'ın kendisi (modül başına
piksel) A ile birebir aynı kalıyor, sadece toplam etiket (yama dahil) daha
büyük çıkıyor.

Kullanım: python3 tests/device/build_print_test_docx_v2.py
Çıktı: tests/device/Baski_Testi_v2_ABC.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image

REGEN_ROOT = Path("/tmp/kalibrasyon_test_2")
OUT = Path(__file__).parent / "Baski_Testi_v2_ABC.docx"

METHODS = ["A", "B", "C"]
STATES = [("state_fresh", "Taze"), ("state_transition", "Geçiş"), ("state_spoiled", "Bozuk")]

BASE_CM = 5.0  # A'nın (yamasız) fiziksel boyutu

# GERÇEK BASKI HATASI (27 Eylül): varsayılan Word sayfa kenar boşluklarıyla
# (1.25in = 3.175cm her yanda, Letter 21.59cm genişlik) kullanılabilir genişlik
# sadece 15.24cm; 3 sütuna bölününce sütun başına ~5.08cm düşüyor. B/C'nin
# gömülen boyutu (~5.55cm, aşağıdaki hesaplamayla) bundan BÜYÜK — sütuna
# sığmayıp kenarları kırpılıyordu (kullanıcı gerçek baskıda fark etti). Düzeltme:
# kenar boşluklarını daraltıp sütun genişliğini açıkça B/C'nin en büyük gömme
# boyutundan belirgin ölçüde geniş tutuyoruz.
MARGIN_CM = 1.27  # 0.5in
COLUMN_CM = 6.3  # 5.55cm'lik en büyük görsel için bolca pay bırakır


def _embed_size_cm(letter: str) -> float:
    """B/C'nin canvas'ı (yama şeridi yüzünden) A'dan büyük — aynı oranda
    büyütülmüş bir fiziksel boyutta gömülürse QR'ın kendi modül boyutu A
    ile birebir eşleşir (yama sadece dışarıda ekstra yer kaplar)."""
    a_png = next((REGEN_ROOT / "A").glob("*.png"))
    this_png = next((REGEN_ROOT / letter).glob("*.png"))
    a_w = Image.open(a_png).size[0]
    this_w = Image.open(this_png).size[0]
    return BASE_CM * this_w / a_w


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(MARGIN_CM)
    section.right_margin = Cm(MARGIN_CM)

    for letter in METHODS:
        folder = REGEN_ROOT / letter
        size = Cm(_embed_size_cm(letter))

        table = doc.add_table(rows=2, cols=3)
        table.autofit = False
        for col, (state_key, state_label) in enumerate(STATES):
            img_path = next(folder.glob(f"*.{state_key}.png"))

            cap_cell = table.cell(0, col)
            cap_cell.width = Cm(COLUMN_CM)
            p = cap_cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"{letter} · {state_label}")
            run.bold = True
            run.font.size = Pt(9)

            img_cell = table.cell(1, col)
            img_cell.width = Cm(COLUMN_CM)
            img_p = img_cell.paragraphs[0]
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.add_run().add_picture(str(img_path), width=size, height=size)

        # python-docx sütun genişliğini bazen sadece hücre bazında değil,
        # tblGrid seviyesinde de tutarlı istiyor — ikisini birden set etmek
        # Word/LibreOffice arasında en güvenilir sonucu veriyor.
        for column in table.columns:
            column.width = Cm(COLUMN_CM)

        doc.add_paragraph()

    doc.save(OUT)
    print(f"Yazıldı: {OUT}")


if __name__ == "__main__":
    main()
