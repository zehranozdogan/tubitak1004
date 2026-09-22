"""B/C etiketlerini 5.5cm boyutunda Word belgesine gömer — A ile aynı modül
boyutuna denk gelsin diye (yama şeridi kenar boşluğunu 4->8 modüle
çıkarıyor, bkz. tests/device/results_2026-09-22.md "yakın çekim" bölümü).

A hâlâ 5x5cm basılı kalıyor (değişmedi) — sadece B/C büyütülüyor.

Kullanım: python3 tests/device/build_print_test_docx_bc_55.py
Çıktı: tests/device/Baski_Testi_BC_5.5cm.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

REGEN_ROOT = Path("/tmp/print_test_regen")
OUT = Path(__file__).parent / "Baski_Testi_BC_5.5cm.docx"

METHODS = [
    ("B", "Yöntem B — white_gray_black", "QR-içi beyaz/siyah + kenarda 1 GRİ referans yaması."),
    ("C", "Yöntem C — multicolor_patch", "QR-içi beyaz/siyah + kenarda gri/kırmızı/yeşil/mavi referans yaması şeridi."),
]
STATES = [("state_fresh", "TAZE"), ("state_transition", "GEÇİŞ"), ("state_spoiled", "BOZUK")]

IMG_SIZE = Cm(5.5)


def add_caption(cell, text: str, *, bold: bool = False, size: int = 10):
    p = cell.paragraphs[0] if not cell.paragraphs[0].runs else cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p


def main() -> None:
    doc = Document()

    title = doc.add_heading("Baskı Testi Etiketleri — B/C (5.5cm, büyütülmüş)", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph()
    intro.add_run(
        "Bu belge SADECE B ve C içindir, 5.5cm boyutunda (A hâlâ 5cm'de kalıyor, "
        "değişmedi). Neden büyütüldü: B/C'nin yama şeridi kenar boşluğunu "
        "4'ten 8 modüle çıkarıyor (canvas 730px->810px, ~%11 büyük); aynı 5cm'e "
        "sığdırılınca QR modülleri A'dakinden ~%10 daha küçük fiziksel boyutta "
        "basılıyordu. 5.5cm'de modül boyutu A ile eşitleniyor.\n"
    ).font.size = Pt(10)
    intro.add_run(
        "Yazdırırken yazıcı diyaloğunda %100 ölçek / gerçek boyut seç, \"kağıda "
        "sığdır\" DEĞİL."
    ).font.size = Pt(10)

    for letter, heading, desc in METHODS:
        folder = REGEN_ROOT / letter
        doc.add_heading(heading, level=2)
        p = doc.add_paragraph(desc)
        p.runs[0].font.size = Pt(9)
        p.runs[0].italic = True

        table = doc.add_table(rows=2, cols=3)
        table.autofit = False

        for col, (state_key, state_label) in enumerate(STATES):
            img_path = next(folder.glob(f"*.{state_key}.png"))
            cap_cell = table.cell(0, col)
            add_caption(cap_cell, f"{letter} · {state_label}", bold=True)

            img_cell = table.cell(1, col)
            img_paragraph = img_cell.paragraphs[0]
            img_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = img_paragraph.add_run()
            run.add_picture(str(img_path), width=IMG_SIZE, height=IMG_SIZE)

        doc.add_paragraph()

    doc.save(OUT)
    print(f"Yazıldı: {OUT}")


if __name__ == "__main__":
    main()
