"""QR tasarım seçeneklerini GERÇEK BOYUTTA basılabilir bir Word dosyasına
koyar — kimya tarafı cetvelle ölçüp "bu blok damla için yeterli mi"
diyebilsin diye (13 Ekim).

Önce seçenekleri üret:  python3 tests/device/build_qr_options.py
Sonra:                  python3 tests/device/build_qr_options_docx.py
"""

from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

KAYNAK = Path(os.environ["HOME"]) / "Desktop" / "QR_Secenekleri"
CIKTI = Path(__file__).parent / "QR_Tasarim_Secenekleri.docx"

# (damla mm, etiket cm, QR matris) -> pratik kombinasyonlar
SECENEKLER = [
    (2.0, 5.0, "65x65"), (2.0, 7.0, "65x65"),
    (3.0, 5.0, "65x65"), (3.0, 7.0, "65x65"),
    (4.0, 5.0, "65x65"), (4.0, 7.0, "65x65"),
]


def main() -> None:
    doc = Document()
    s = doc.sections[0]
    s.left_margin = s.right_margin = Cm(1.27)
    s.top_margin = s.bottom_margin = Cm(1.27)

    b = doc.add_paragraph().add_run("QR tasarım seçenekleri — manuel damlatma için")
    b.bold = True
    b.font.size = Pt(14)

    a = doc.add_paragraph().add_run(
        "Bu sayfadaki QR'lar GERÇEK BOYUTTA basılmıştır (yazdırırken "
        "'sayfaya sığdır'/ölçekleme KAPALI olmalı). "
        "SOLDAKİ görsel basılacak olandır: reaktif alanlar BOŞ bırakılır "
        "(beyaz kâğıt), yalnızca bir modüllük çerçeveleri basılır — çerçeve, "
        "operatörün nereye damlatacağını görmesi içindir. İndikatör beyaz "
        "kâğıdın üzerine damlatılmalı; altına renk basılırsa ölçüm bozulur. "
        "SAĞDAKİ görsel yalnızca indikatör uygulandıktan sonra nasıl "
        "görüneceğini gösterir, basılmaz. Çerçevelerin içini cetvelle ölçüp "
        "damlanın sığıp sığmadığına bakılabilir.")
    a.font.size = Pt(9)

    for damla, etiket, matris in SECENEKLER:
        ad = f"{matris}_{etiket:.0f}cm_damla{damla:.0f}mm"
        baski = KAYNAK / f"{ad}_baski.png"
        ornek = KAYNAK / f"{ad}_ornek.png"
        if not baski.exists():
            continue
        doc.add_paragraph()
        bas = doc.add_paragraph().add_run(
            f"Etiket {etiket:.0f} cm · reaktif alan ≈ {damla:.0f} mm hedefli")
        bas.bold = True
        bas.font.size = Pt(11)
        tablo = doc.add_table(rows=2, cols=2)
        for sutun, (yol, baslik) in enumerate([
            (baski, "BASILACAK OLAN — çerçeveli kareler boş bırakılır,\nindikatör oraya damlatılır"),
            (ornek, "ÖRNEK — indikatör uygulanıp renk değiştikten sonraki görünüm\n(bu BASILMAZ)"),
        ]):
            h = tablo.cell(0, sutun).paragraphs[0]
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hr = h.add_run(baslik)
            hr.bold = True
            hr.font.size = Pt(8)
            g = tablo.cell(1, sutun).paragraphs[0]
            g.alignment = WD_ALIGN_PARAGRAPH.CENTER
            g.add_run().add_picture(str(yol), width=Cm(etiket), height=Cm(etiket))
        alt = doc.add_paragraph()
        alt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = alt.add_run(f"{ad}  —  çerçevelerin İÇİ ≈ {damla:.0f} mm "
                        f"(damlanın sığması gereken alan)")
        r.font.size = Pt(8)

    doc.save(CIKTI)
    print(f"Yazıldı: {CIKTI}")


if __name__ == "__main__":
    main()
