"""Kağıt baskı testi (3. tur, 4 Ekim) — A/B/C kalibrasyon karşılaştırması.

ÖNCEKİ TURLARDAN FARKI — ikisi de o turları geçersiz kılan hatalardı:

1. Etiketler DART motoruyla üretiliyor (v2 Python'la üretiyordu). İki motor
   aynı payload için farklı QR deseni üretiyor; Python etiketi Dart
   okuyucuyla okunamıyor (bkz. docs/decisions/0004).
2. Her etiket KENDİ kalibrasyon profiline işaret ediyor
   (GENIPIN_PUTRESIN_v2_A/B/C). v2'de üçü de aynı profili gösterdiği için
   okuyucu üçünü de white_black (A) ile okuyordu — yani B ve C hiç test
   edilmemişti.

ÖLÇÜ: B/C'nin yamaları QR'ın DIŞINA taştığı için toplam görsel A'dan
büyük (810 vs 730 px). Üçü de aynı fiziksel boyutta basılırsa QR'ların
MODÜL boyutu farklı olur ve karşılaştırma bozulur. Bu yüzden B/C oranlı
şekilde büyük basılıyor (5.55 cm) — QR'ın kendisi üçünde de aynı.

SAYFA DÜZENİ: varsayılan Word kenar boşluklarıyla sütun genişliği 5.08cm
kalıyor ve 5.55cm'lik görseller kırpılıyordu (22 Eylül turunda gerçekten
yaşandı). Kenar boşlukları daraltılıp sütun genişliği açıkça veriliyor.

Önce etiketleri üret:
  cd packages_dart/label_export && dart run tool/generate_calibration_labels.dart
Sonra:
  python3 tests/device/build_print_test_docx_v3.py
"""

from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image

KAYNAK = Path(os.environ["HOME"]) / "Desktop" / "QR_Test_Dart"
CIKTI = Path(__file__).parent / "Baski_Testi_v3_ABC.docx"

YONTEMLER = ["A", "B", "C"]
DURUMLAR = [("state_fresh", "Taze"), ("state_transition", "Geçiş"), ("state_spoiled", "Bozuk")]

TABAN_CM = 5.0        # A'nın (yamasız) fiziksel boyutu
KENAR_CM = 1.27       # 0.5 inç
SUTUN_CM = 6.3        # en büyük görsel (5.55cm) için bolca pay


def _png(harf: str, durum: str) -> Path:
    return KAYNAK / harf / f"KALIB-{harf}_QR_SENSOR_v4.{durum}.png"


def _gomme_boyutu_cm(harf: str) -> float:
    """QR'ın modül boyutu üçünde de aynı kalsın diye oranlı büyütme."""
    a_px = Image.open(_png("A", "state_fresh")).size[0]
    bu_px = Image.open(_png(harf, "state_fresh")).size[0]
    return TABAN_CM * bu_px / a_px


def main() -> None:
    eksik = [str(_png(h, d)) for h in YONTEMLER for d, _ in DURUMLAR if not _png(h, d).exists()]
    if eksik:
        raise SystemExit(
            "Etiket bulunamadı:\n  " + "\n  ".join(eksik)
            + "\n\nÖnce üret:\n  cd packages_dart/label_export"
              " && dart run tool/generate_calibration_labels.dart"
        )

    doc = Document()
    bolum = doc.sections[0]
    bolum.left_margin = Cm(KENAR_CM)
    bolum.right_margin = Cm(KENAR_CM)

    for harf in YONTEMLER:
        boyut = Cm(_gomme_boyutu_cm(harf))
        tablo = doc.add_table(rows=2, cols=3)
        tablo.autofit = False
        for sutun, (durum, etiket) in enumerate(DURUMLAR):
            baslik = tablo.cell(0, sutun)
            baslik.width = Cm(SUTUN_CM)
            p = baslik.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"{harf} · {etiket}")
            run.bold = True
            run.font.size = Pt(9)

            gorsel = tablo.cell(1, sutun)
            gorsel.width = Cm(SUTUN_CM)
            gp = gorsel.paragraphs[0]
            gp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            gp.add_run().add_picture(str(_png(harf, durum)), width=boyut, height=boyut)

        for sutun in tablo.columns:
            sutun.width = Cm(SUTUN_CM)
        doc.add_paragraph()

    doc.save(CIKTI)
    print(f"Yazıldı: {CIKTI}")
    for harf in YONTEMLER:
        print(f"  {harf}: {_gomme_boyutu_cm(harf):.2f} cm")


if __name__ == "__main__":
    main()
